#!/usr/bin/env python3
"""Structural check of an OOXML package (.pptx / .docx) written by Daftar.

1. zip integrity; 2. every XML / .rels part parses (xml.dom.minidom); 3. every part has a content type and every
Override names an existing part; 4. every internal relationship target exists; 5. format specific checks;
6. optional: open with python-pptx / python-docx when installed.
"""
import posixpath
import re
import sys
import zipfile
from xml.dom import minidom

ILLEGAL = re.compile('[\x00-\x08\x0b\x0c\x0e-\x1f]')


def rels_name_for(part):
    d, b = posixpath.split(part)
    return posixpath.join(d, '_rels', b + '.rels')


def check(path):
    errs = []
    z = zipfile.ZipFile(path)
    bad = z.testzip()
    if bad:
        errs.append('corrupt zip member ' + bad)
    names = [n for n in z.namelist() if not n.endswith('/')]
    nameset = set(names)
    doms = {}
    for n in names:
        if n.endswith('.xml') or n.endswith('.rels'):
            raw = z.read(n)
            txt = raw.decode('utf-8')
            if ILLEGAL.search(txt):
                errs.append('illegal XML 1.0 control char in ' + n)
            try:
                doms[n] = minidom.parseString(raw)
            except Exception as e:  # noqa
                errs.append('XML parse error in %s: %s' % (n, e))
    ct = doms.get('[Content_Types].xml')
    if ct is None:
        errs.append('missing [Content_Types].xml')
        return errs, names
    defaults = {e.getAttribute('Extension').lower(): e.getAttribute('ContentType') for e in ct.getElementsByTagName('Default')}
    overrides = {e.getAttribute('PartName'): e.getAttribute('ContentType') for e in ct.getElementsByTagName('Override')}
    for p in overrides:
        if p.lstrip('/') not in nameset:
            errs.append('Override for missing part ' + p)
    for n in names:
        if n == '[Content_Types].xml':
            continue
        if '/' + n not in overrides and n.rsplit('.', 1)[-1].lower() not in defaults:
            errs.append('no content type for ' + n)
    # relationships
    for n in names:
        if not n.endswith('.rels'):
            continue
        src_dir = posixpath.dirname(posixpath.dirname(n))
        ids = set()
        for r in doms[n].getElementsByTagName('Relationship'):
            rid = r.getAttribute('Id')
            if rid in ids:
                errs.append('duplicate rel id %s in %s' % (rid, n))
            ids.add(rid)
            if r.getAttribute('TargetMode') == 'External':
                continue
            t = r.getAttribute('Target')
            full = t.lstrip('/') if t.startswith('/') else posixpath.normpath(posixpath.join(src_dir, t))
            if full not in nameset:
                errs.append('%s: %s -> missing %s' % (n, rid, full))
    root_rels = doms.get('_rels/.rels')
    if root_rels is None:
        errs.append('missing _rels/.rels')
    lower = path.lower()
    if lower.endswith('.pptx'):
        errs += check_pptx(z, doms, nameset)
    elif lower.endswith('.docx'):
        errs += check_docx(doms)
    return errs, names


def rel_map(doms, part):
    d = doms.get(rels_name_for(part))
    if d is None:
        return {}
    base = posixpath.dirname(part)
    return {r.getAttribute('Id'): posixpath.normpath(posixpath.join(base, r.getAttribute('Target')))
            for r in d.getElementsByTagName('Relationship')}


def check_pptx(z, doms, nameset):
    errs = []
    pres = doms.get('ppt/presentation.xml')
    if pres is None:
        return ['missing ppt/presentation.xml']
    prels = rel_map(doms, 'ppt/presentation.xml')
    sz = pres.getElementsByTagName('p:sldSz')[0]
    cx, cy = int(sz.getAttribute('cx')), int(sz.getAttribute('cy'))
    for v in (cx, cy):
        if not 914400 <= v <= 51206400:
            errs.append('sldSz out of range %d' % v)
    ids = []
    for s in pres.getElementsByTagName('p:sldId'):
        sid = int(s.getAttribute('id'))
        ids.append(sid)
        if not 256 <= sid < 2147483648:
            errs.append('bad sldId %d' % sid)
        tgt = prels.get(s.getAttribute('r:id'))
        if tgt is None or not tgt.startswith('ppt/slides/slide'):
            errs.append('sldId %d has no slide rel' % sid)
            continue
        srels = rel_map(doms, tgt)
        sd = doms[tgt]
        for b in sd.getElementsByTagName('a:blip'):
            if b.getAttribute('r:embed') not in srels:
                errs.append('%s: blip rel %s missing' % (tgt, b.getAttribute('r:embed')))
        for off in sd.getElementsByTagName('a:off'):
            pass
        for p in sd.getElementsByTagName('p:pic'):
            xfrm = p.getElementsByTagName('a:xfrm')[0]
            o = xfrm.getElementsByTagName('a:off')[0]
            e = xfrm.getElementsByTagName('a:ext')[0]
            x, y = int(o.getAttribute('x')), int(o.getAttribute('y'))
            w, h = int(e.getAttribute('cx')), int(e.getAttribute('cy'))
            if x < 0 or y < 0 or x + w > cx or y + h > cy:
                errs.append('%s: picture outside slide' % tgt)
    if len(set(ids)) != len(ids):
        errs.append('duplicate slide ids')
    mids = [int(m.getAttribute('id')) for m in pres.getElementsByTagName('p:sldMasterId')]
    for m in pres.getElementsByTagName('p:sldMasterId'):
        master = prels.get(m.getAttribute('r:id'))
        md = doms.get(master)
        if md is None:
            errs.append('master missing')
            continue
        mrels = rel_map(doms, master)
        for l in md.getElementsByTagName('p:sldLayoutId'):
            mids.append(int(l.getAttribute('id')))
            if l.getAttribute('r:id') not in mrels:
                errs.append('layout rel missing')
        if not md.getElementsByTagName('p:clrMap'):
            errs.append('master without clrMap')
    if any(i < 2147483648 for i in mids) or len(set(mids)) != len(mids):
        errs.append('master/layout ids invalid %s' % mids)
    theme = doms.get('ppt/theme/theme1.xml')
    if theme is None:
        errs.append('no theme')
    else:
        def count(lst, child):
            el = theme.getElementsByTagName(lst)[0]
            return len([c for c in el.childNodes if c.nodeType == 1])
        for lst in ('a:fillStyleLst', 'a:lnStyleLst', 'a:effectStyleLst', 'a:bgFillStyleLst'):
            if count(lst, None) < 3:
                errs.append('theme %s has < 3 styles' % lst)
        if len(theme.getElementsByTagName('a:clrScheme')[0].childNodes) != 12:
            errs.append('clrScheme must have 12 colours')
    return errs


def check_docx(doms):
    errs = []
    doc = doms.get('word/document.xml')
    if doc is None:
        return ['missing word/document.xml']
    if not doc.getElementsByTagName('w:body'):
        errs.append('document.xml has no w:body')
    return errs


def library_open(path):
    try:
        if path.lower().endswith('.pptx'):
            import pptx
            p = pptx.Presentation(path)
            return 'python-pptx: %d slides, %dx%d EMU, pics/slide=%s' % (
                len(p.slides), p.slide_width, p.slide_height,
                [sum(1 for s in sl.shapes if s.shape_type == 13) for sl in p.slides])
        if path.lower().endswith('.docx'):
            import docx
            d = docx.Document(path)
            brs = sum(1 for p in d.paragraphs for r in p.runs for b in r._r.findall(
                '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}br') if b.get(
                '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type') == 'page')
            return 'python-docx: %d paragraphs, %d page breaks, first=%r' % (len(d.paragraphs), brs, d.paragraphs[0].text[:40] if d.paragraphs else '')
    except ImportError:
        return 'library not installed (structural checks only)'
    except Exception as e:  # noqa
        return 'LIBRARY FAILED: %r' % e
    return ''


if __name__ == '__main__':
    rc = 0
    for f in sys.argv[1:]:
        errs, names = check(f)
        print('%s: %d parts, %s' % (f, len(names), 'OK' if not errs else 'FAIL'))
        for e in errs:
            print('   - ' + e)
        print('   ' + library_open(f))
        rc |= bool(errs)
    sys.exit(rc)
