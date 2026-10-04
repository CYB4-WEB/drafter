"""Builds a realistic student library for emulator testing (docs/seed/Library)."""
import json, os, shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from docx import Document
from docx.shared import Pt as DPt
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

ROOT = os.path.join(os.path.dirname(__file__), "..", "docs", "seed", "Library")
shutil.rmtree(ROOT, ignore_errors=True)

def folder(path, color, icon, desc=""):
    os.makedirs(path, exist_ok=True)
    json.dump({"color": color, "icon": icon, "desc": desc}, open(os.path.join(path, ".meta.json"), "w"))

def course(name, color, icon, desc):
    base = os.path.join(ROOT, name)
    folder(base, color, icon, desc)
    for sub, ic in [("Lectures", "lecture"), ("Seminars", "seminar"), ("Labs", "lab")]:
        folder(os.path.join(base, sub), color, ic)
    return base

def pdf(path, title, pages):
    c = canvas.Canvas(path, pagesize=A4)
    w, h = A4
    for i, (head, lines) in enumerate(pages):
        c.setFont("Helvetica-Bold", 22); c.drawString(60, h - 80, head)
        c.setFont("Helvetica", 12)
        y = h - 120
        for ln in lines:
            c.drawString(60, y, ln); y -= 20
        if i == 0:
            # simple OSI table drawing
            for k, layer in enumerate(["7 Application", "6 Presentation", "5 Session", "4 Transport", "3 Network", "2 Data Link", "1 Physical"]):
                c.rect(200, h - 420 - k * 28, 200, 28); c.drawString(215, h - 410 - k * 28, layer)
        c.setFont("Helvetica", 9); c.drawString(w - 90, 30, f"{title} - {i+1}")
        c.showPage()
    c.save()

def pptx(path, rtl=False):
    p = Presentation()
    p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
    s = p.slides.add_slide(p.slide_layouts[0])
    s.shapes.title.text = "مقدمة في الشبكات" if rtl else "Routing Fundamentals"
    s.placeholders[1].text = "المحاضرة ٣ — د. أحمد" if rtl else "Lecture 3 · Network & CCNA"
    s.notes_slide.notes_text_frame.text = "رحب بالطلاب واذكر موعد الاختبار" if rtl else "Welcome students. Remind them the quiz is next Tuesday."
    s2 = p.slides.add_slide(p.slide_layouts[1])
    s2.shapes.title.text = "أنواع التوجيه" if rtl else "Static vs Dynamic Routing"
    tf = s2.placeholders[1].text_frame
    items = ["التوجيه الثابت: يضبطه المسؤول يدويًا", "التوجيه الديناميكي: OSPF و EIGRP", "المقياس والمسافة الإدارية"] if rtl else \
        ["Static routes are configured manually", "Dynamic protocols: OSPF, EIGRP, BGP", "Administrative distance decides between sources", "Metric picks the best path"]
    tf.text = items[0]
    for it in items[1:]:
        para = tf.add_paragraph(); para.text = it
    s2.notes_slide.notes_text_frame.text = "Explain AD values: connected 0, static 1, OSPF 110."
    s3 = p.slides.add_slide(p.slide_layouts[5])
    s3.shapes.title.text = "Administrative distance"
    rows = [("Source", "AD"), ("Connected", "0"), ("Static", "1"), ("EIGRP", "90"), ("OSPF", "110"), ("RIP", "120")]
    tbl = s3.shapes.add_table(len(rows), 2, Inches(1.5), Inches(1.8), Inches(6), Inches(3.6)).table
    for r, (a, b) in enumerate(rows):
        tbl.cell(r, 0).text = a; tbl.cell(r, 1).text = b
    box = s3.shapes.add_shape(1, Inches(8.3), Inches(2.2), Inches(4), Inches(1.6))
    box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0x3B, 0x82, 0xF6)
    box.text_frame.text = "Lower AD wins"
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    box.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    p.save(path)

def docx(path, rtl=False):
    d = Document()
    if rtl:
        d.add_heading("تقرير المختبر ٢: إعداد الشبكة", 0)
        d.add_paragraph("الهدف من هذا المختبر هو إعداد موجهين وربطهما باستخدام بروتوكول OSPF.")
        d.add_heading("الخطوات", 1)
        for t in ["وصّل الموجه الأول بالمحول", "اضبط عناوين IP", "فعّل بروتوكول OSPF في المنطقة 0"]:
            d.add_paragraph(t, style="List Number")
    else:
        d.add_heading("Lab 2: Configuring OSPF", 0)
        p = d.add_paragraph("Objective: ")
        p.add_run("configure two routers").bold = True
        p.add_run(" and verify neighbour adjacency using ").italic = False
        p.add_run("show ip ospf neighbor").italic = True
        d.add_heading("Equipment", 1)
        for t in ["2 × Cisco 4321 routers", "1 × Catalyst 2960 switch", "Console cables"]:
            d.add_paragraph(t, style="List Bullet")
        d.add_heading("Steps", 1)
        for t in ["Cable the topology", "Assign IP addresses", "Enable OSPF area 0", "Verify with ping"]:
            d.add_paragraph(t, style="List Number")
        d.add_heading("Addressing table", 2)
        t = d.add_table(rows=4, cols=3); t.style = "Table Grid"
        for r, row in enumerate([("Device", "Interface", "IP"), ("R1", "G0/0", "10.0.0.1/30"), ("R2", "G0/0", "10.0.0.2/30"), ("R1", "Lo0", "1.1.1.1/32")]):
            for cidx, v in enumerate(row): t.cell(r, cidx).text = v
        d.add_paragraph("Submit your report before Thursday 23:59.")
    d.save(path)

net = course("Network & CCNA", 4, "globe", "CCNA 2 — Dr. Ahmed")
cyber = course("Cyber Security", 6, "code", "SEC301")
prog = course("Programming", 1, "code", "Kotlin & Android")
uni = course("University", 0, "lecture", "")
folder(os.path.join(ROOT, "Personal"), 2, "star")

pdf(os.path.join(net, "Lectures", "Network_Notes.pdf"), "Network Notes", [
    ("1. Network Fundamentals", ["1.1 OSI Model", "Each layer serves the layer above it."]),
    ("1.2 TCP/IP Model", ["Application, Transport, Internet, Network Access", "Encapsulation adds headers at each layer."]),
    ("2. Addressing", ["IPv4: 32-bit, dotted decimal", "Subnetting: borrow host bits", "/30 gives 2 usable hosts"]),
])
pdf(os.path.join(cyber, "Lectures", "OWASP_Top10.pdf"), "OWASP", [
    ("OWASP Top 10", ["A01 Broken Access Control", "A02 Cryptographic Failures", "A03 Injection"]),
    ("Injection", ["Never build SQL from strings", "Use parameterised queries"]),
])
pptx(os.path.join(net, "Lectures", "Lecture 3 - Routing.pptx"))
pptx(os.path.join(net, "Seminars", "مقدمة في الشبكات.pptx"), rtl=True)
docx(os.path.join(net, "Labs", "Lab 2 - OSPF.docx"))
docx(os.path.join(uni, "Labs", "تقرير المختبر.docx"), rtl=True)
docx(os.path.join(cyber, "Seminars", "Pentest Checklist.docx"))
pptx(os.path.join(prog, "Lectures", "Kotlin Basics.pptx"))
print("seed ok")
