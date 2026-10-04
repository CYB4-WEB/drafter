import io, zipfile, os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image, ImageDraw

OUT = r"C:\Users\am204\AndroidStudioProjects\Daftar\docs\testdata"
os.makedirs(OUT, exist_ok=True)


def png(w, h, c1, c2):
    im = Image.new("RGB", (w, h), c1)
    d = ImageDraw.Draw(im)
    for i in range(0, w, 40):
        d.rectangle([i, 0, i + 19, h], fill=c2)
    d.ellipse([w // 4, h // 4, 3 * w // 4, 3 * h // 4], fill=(255, 255, 255))
    b = io.BytesIO(); im.save(b, "PNG"); b.seek(0); return b


def rtl(par):
    pPr = par._p.get_or_add_pPr(); pPr.set("rtl", "1"); pPr.set("algn", "r")


def build(path, arabic):
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    L = prs.slide_layouts
    # 1 title slide
    s = prs.slides.add_slide(L[0])
    s.shapes.title.text = "مقدمة في أمن المعلومات" if arabic else "Introduction to Cyber Security"
    s.placeholders[1].text = "المحاضرة 1 — Dr. Ahmed" if arabic else "Lecture 1 — Dr. Smith"
    if arabic:
        for ph in (s.shapes.title, s.placeholders[1]):
            for p in ph.text_frame.paragraphs:
                p._p.get_or_add_pPr().set("rtl", "1")
    s.notes_slide.notes_text_frame.text = "رحّب بالطلاب وقدّم المنهج." if arabic else "Welcome students and present the syllabus."
    # 2 bullets (legacy comments)
    s = prs.slides.add_slide(L[1])
    s.shapes.title.text = "المفاهيم الأساسية" if arabic else "Core concepts"
    tf = s.placeholders[1].text_frame
    items = ([("السرية Confidentiality", 0), ("النزاهة Integrity والتحقق", 1), ("التوفر Availability", 0), ("مثال: هجوم DDoS على الخادم", 1)]
             if arabic else
             [("Confidentiality", 0), ("Encryption at rest and in transit", 1), ("Integrity", 0), ("Hashes, MACs, signatures", 1), ("Availability", 0)])
    tf.text = items[0][0]
    if arabic: rtl(tf.paragraphs[0])
    for t, lvl in items[1:]:
        p = tf.add_paragraph(); p.text = t; p.level = lvl
        if arabic: rtl(p)
    s.notes_slide.notes_text_frame.text = ("اشرح مثلث CIA بأمثلة.\nاسأل الطلاب عن أمثلة." if arabic
                                           else "Explain the CIA triad with examples.\nAsk students for real incidents.")
    # 3 image + caption + slide background (modern comments)
    s = prs.slides.add_slide(L[5])
    s.shapes.title.text = "مخطط الشبكة" if arabic else "Network diagram"
    s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor(0xF4, 0xF1, 0xE8)
    s.shapes.add_picture(png(800, 450, (59, 130, 246), (30, 64, 175)), Inches(1), Inches(1.6), Inches(7), Inches(3.94))
    tb = s.shapes.add_textbox(Inches(8.4), Inches(2), Inches(4.2), Inches(2))
    tb.text_frame.word_wrap = True
    tb.text_frame.text = "الشكل 1: جدار حماية firewall بين الشبكتين" if arabic else "Figure 1: a firewall between two networks"
    if arabic: rtl(tb.text_frame.paragraphs[0])
    tb.text_frame.paragraphs[0].runs[0].font.size = Pt(20)
    tb.text_frame.paragraphs[0].runs[0].font.italic = True
    # 4 table
    s = prs.slides.add_slide(L[5])
    s.shapes.title.text = "مقارنة الخوارزميات" if arabic else "Algorithm comparison"
    rows = ([("الخوارزمية", "النوع", "طول المفتاح"), ("AES", "متماثل", "128/256"), ("RSA", "غير متماثل", "2048+"), ("SHA-256", "تجزئة", "—")]
            if arabic else
            [("Algorithm", "Type", "Key length"), ("AES", "Symmetric", "128/256"), ("RSA", "Asymmetric", "2048+"), ("SHA-256", "Hash", "—")])
    t = s.shapes.add_table(4, 3, Inches(1), Inches(1.8), Inches(11), Inches(3)).table
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
            if arabic: rtl(t.cell(r, c).text_frame.paragraphs[0])
    # 5 shapes + group + connector
    s = prs.slides.add_slide(L[5])
    s.shapes.title.text = "الأشكال" if arabic else "Shapes and arrows"
    x = Inches(0.8)
    for shp, label in [(MSO_SHAPE.ROUNDED_RECTANGLE, "Plan"), (MSO_SHAPE.CHEVRON, "Do"), (MSO_SHAPE.OVAL, "Check"),
                       (MSO_SHAPE.RIGHT_ARROW, "Act"), (MSO_SHAPE.ISOSCELES_TRIANGLE, "!")]:
        sh = s.shapes.add_shape(shp, x, Inches(2), Inches(2.2), Inches(1.4)); sh.text = label; x += Inches(2.45)
    grp = s.shapes.add_group_shape()
    a = grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2), Inches(4.6), Inches(3), Inches(1)); a.text = "Group A"
    b = grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6), Inches(4.6), Inches(3), Inches(1)); b.text = "Group B"
    b.fill.solid(); b.fill.fore_color.rgb = RGBColor(0x4C, 0xC3, 0x8A)
    ln = s.shapes.add_connector(1, Inches(5), Inches(5.1), Inches(6), Inches(5.1))
    ln.line.width = Pt(2)
    prs.save(path)


def inject_comments(path, arabic):
    zin = zipfile.ZipFile(path); files = {n: zin.read(n) for n in zin.namelist()}; zin.close()
    ct = files["[Content_Types].xml"].decode()
    ct = ct.replace("</Types>",
                    '<Override PartName="/ppt/comments/comment1.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.comments+xml"/>'
                    '<Override PartName="/ppt/commentAuthors.xml" ContentType="application/vnd.openxmlformats-officedocument.presentationml.commentAuthors+xml"/>'
                    '<Override PartName="/ppt/comments/modernComment_1.xml" ContentType="application/vnd.ms-powerpoint.comments+xml"/>'
                    '<Override PartName="/ppt/authors.xml" ContentType="application/vnd.ms-powerpoint.authors+xml"/></Types>')
    files["[Content_Types].xml"] = ct.encode()
    P = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    files["ppt/commentAuthors.xml"] = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><p:cmAuthorLst {P}>'
        '<p:cmAuthor id="0" name="Dr. Smith" initials="DS" lastIdx="2" clrIdx="0"/>'
        '<p:cmAuthor id="1" name="Sara Ali" initials="SA" lastIdx="1" clrIdx="1"/></p:cmAuthorLst>').encode()
    t1 = "راجع هذا التعريف قبل الامتحان" if arabic else "Review this definition before the exam"
    t2 = "شكرًا، سأفعل" if arabic else "Thanks, will do"
    files["ppt/comments/comment1.xml"] = (
        f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><p:cmLst {P} xmlns:p15="http://schemas.microsoft.com/office/powerpoint/2012/main">'
        f'<p:cm authorId="0" dt="2026-09-20T10:15:00.000" idx="1"><p:pos x="10" y="10"/><p:text>{t1}</p:text></p:cm>'
        f'<p:cm authorId="1" dt="2026-09-21T08:00:00.000" idx="1"><p:pos x="10" y="10"/><p:text>{t2}</p:text>'
        '<p:extLst><p:ext uri="{C676402C-5697-4E1C-873F-D02D1690AC5C}"><p15:threadingInfo timeZoneBias="0">'
        '<p15:parentCm authorId="0" idx="1"/></p15:threadingInfo></p:ext></p:extLst></p:cm></p:cmLst>').encode()
    files["ppt/authors.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><p188:authorLst xmlns:p188="http://schemas.microsoft.com/office/powerpoint/2018/8/main">'
        '<p188:author id="{11111111-AAAA-BBBB-CCCC-000000000001}" name="Prof. Karim" initials="PK" userId="Prof. Karim" providerId="None"/>'
        '<p188:author id="{11111111-AAAA-BBBB-CCCC-000000000002}" name="Lina" initials="L" userId="Lina" providerId="None"/></p188:authorLst>').encode()
    m1 = "هل يمكن إضافة مثال عن AES؟" if arabic else "Can we add an AES example here?"
    m2 = "أضفته في الشريحة 4" if arabic else "Added it on slide 4"
    files["ppt/comments/modernComment_1.xml"] = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?><p188:cmLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'xmlns:p188="http://schemas.microsoft.com/office/powerpoint/2018/8/main">'
        '<p188:cm id="{22222222-0000-0000-0000-000000000001}" authorId="{11111111-AAAA-BBBB-CCCC-000000000001}" created="2026-09-22T12:00:00.000">'
        '<p188:replyLst><p188:reply id="{33333333-0000-0000-0000-000000000001}" authorId="{11111111-AAAA-BBBB-CCCC-000000000002}" created="2026-09-22T13:30:00.000">'
        f'<p188:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US"/><a:t>{m2}</a:t></a:r></a:p></p188:txBody></p188:reply></p188:replyLst>'
        f'<p188:txBody><a:bodyPr/><a:lstStyle/><a:p><a:r><a:rPr lang="en-US"/><a:t>{m1}</a:t></a:r></a:p></p188:txBody></p188:cm></p188:cmLst>').encode()

    def add_rel(rels_path, rid, typ, target):
        x = files[rels_path].decode()
        files[rels_path] = x.replace("</Relationships>", f'<Relationship Id="{rid}" Type="{typ}" Target="{target}"/></Relationships>').encode()

    add_rel("ppt/slides/_rels/slide2.xml.rels", "rIdC1", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments", "../comments/comment1.xml")
    add_rel("ppt/slides/_rels/slide3.xml.rels", "rIdC2", "http://schemas.microsoft.com/office/2018/10/relationships/comments", "../comments/modernComment_1.xml")
    add_rel("ppt/_rels/presentation.xml.rels", "rIdCA", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/commentAuthors", "commentAuthors.xml")
    add_rel("ppt/_rels/presentation.xml.rels", "rIdAU", "http://schemas.microsoft.com/office/2018/10/relationships/authors", "authors.xml")
    zout = zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED)
    for n, d in files.items():
        zout.writestr(n, d)
    zout.close()


for name, ar in [("lecture_en.pptx", False), ("lecture_ar.pptx", True)]:
    p = os.path.join(OUT, name); build(p, ar); inject_comments(p, ar); print("wrote", p)
open(os.path.join(OUT, "corrupt.pptx"), "wb").write(b"PK\x03\x04 this is not really a zip file")
open(os.path.join(OUT, "legacy.ppt"), "wb").write(bytes([0xD0, 0xCF, 0x11, 0xE0, 0xA1, 0xB1, 0x1A, 0xE1]) + b"\0" * 504)
print("ok")
